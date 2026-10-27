"""Spreadsheet parsers (XLSX via openpyxl, legacy XLS via xlrd) and CSV.

Spreadsheets are never flattened into one blob: each sheet is split into table
regions that keep sheet names, cell coordinates, header rows, row numbers,
number formats and formula text (formulas are recorded, never evaluated;
macros are never executed - openpyxl/xlrd do not run VBA).
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

from app.config import get_settings
from app.extraction.base import DocumentParser, ExtractedDocument, ExtractionError
from app.extraction.tabular import GridCell, GridRow, find_unsafe_formulas, format_value, regions_to_blocks


class XlsxParser(DocumentParser):
    types = frozenset({"xlsx"})

    def parse(self, path: Path) -> ExtractedDocument:
        import openpyxl

        settings = get_settings()
        out = ExtractedDocument(detected_type="xlsx")
        try:
            with zipfile.ZipFile(path) as zf:
                if any(n.lower().endswith("vbaproject.bin") for n in zf.namelist()):
                    out.warnings.append("Workbook contains a VBA macro project; macros were ignored and never executed.")
                if any(n.lower().startswith("xl/externallinks/") for n in zf.namelist()):
                    out.warnings.append("Workbook references external workbooks; external links were not followed.")
            wb_formula = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
            wb_values = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        except Exception as exc:
            raise ExtractionError(f"Could not open workbook: {type(exc).__name__}") from exc

        total_cells = 0
        table_index = 0
        uncalculated = 0
        sheets_info = []
        try:
            for ws_f in wb_formula.worksheets:
                ws_v = wb_values[ws_f.title]
                state = getattr(ws_f, "sheet_state", "visible")
                rows: list[GridRow] = []
                for (r_idx, row_f), row_v in zip(
                    enumerate(ws_f.iter_rows(), start=1), ws_v.iter_rows(values_only=True)
                ):
                    cells: list[GridCell] = []
                    for c_idx, cell in enumerate(row_f):
                        total_cells += 1
                        if total_cells > settings.max_spreadsheet_cells:
                            raise ExtractionError(
                                f"Spreadsheet exceeds the configured cell limit ({settings.max_spreadsheet_cells})"
                            )
                        raw_f = getattr(cell, "value", None)
                        cached = row_v[c_idx] if c_idx < len(row_v) else None
                        number_format = getattr(cell, "number_format", None)
                        formula = raw_f if isinstance(raw_f, str) and raw_f.startswith("=") else None
                        if formula and cached is None:
                            uncalculated += 1
                            display = formula  # show the formula text rather than inventing a value
                        else:
                            display = format_value(cached if formula else raw_f, number_format)
                        cells.append(GridCell(display, cached if formula else raw_f, formula, number_format))
                    actual_row = getattr(row_f[0], "row", r_idx) if row_f else r_idx
                    rows.append(GridRow(number=actual_row or r_idx, cells=cells))
                blocks = regions_to_blocks(rows, sheet=ws_f.title, method="openpyxl", start_table_index=table_index)
                table_index += len(blocks)
                for b in blocks:
                    unsafe = find_unsafe_formulas(b.table.formulas if b.table else {})
                    if unsafe:
                        out.warnings.append(
                            f"Sheet '{ws_f.title}': potentially unsafe formulas at {', '.join(unsafe[:5])} (recorded as text, not evaluated)."
                        )
                out.blocks.extend(blocks)
                sheets_info.append({"name": ws_f.title, "state": state, "tables": len(blocks)})
        finally:
            wb_formula.close()
            wb_values.close()
        if uncalculated:
            out.warnings.append(
                f"{uncalculated} formula cell(s) had no cached value; formula text is shown instead of a computed result."
            )
        out.page_count = len(sheets_info)
        out.metadata["sheets"] = sheets_info
        out.quality = {"cells_scanned": total_cells, "tables": table_index, "uncalculated_formulas": uncalculated}
        return out


class XlsParser(DocumentParser):
    types = frozenset({"xls"})

    def parse(self, path: Path) -> ExtractedDocument:
        import xlrd

        settings = get_settings()
        out = ExtractedDocument(detected_type="xls")
        try:
            book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
        except Exception:
            try:
                book = xlrd.open_workbook(str(path), on_demand=True)
            except Exception as exc:
                raise ExtractionError(f"Could not open XLS workbook: {type(exc).__name__}") from exc
        out.warnings.append("Legacy .xls: formula text is not available; cached values are used.")
        total_cells = 0
        table_index = 0
        sheets_info = []
        try:
            for sheet_idx in range(book.nsheets):
                sh = book.sheet_by_index(sheet_idx)
                total_cells += sh.nrows * sh.ncols
                if total_cells > settings.max_spreadsheet_cells:
                    raise ExtractionError("Spreadsheet exceeds the configured cell limit")
                rows: list[GridRow] = []
                for r in range(sh.nrows):
                    cells = []
                    for c in range(sh.ncols):
                        cell = sh.cell(r, c)
                        fmt = None
                        if book.formatting_info:
                            try:
                                xf = book.xf_list[cell.xf_index]
                                fmt = book.format_map[xf.format_key].format_str
                            except Exception:
                                fmt = None
                        value: object = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            try:
                                value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                            except Exception:
                                pass
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = "#ERR"
                        cells.append(GridCell(format_value(value, fmt), value, None, fmt))
                    rows.append(GridRow(number=r + 1, cells=cells))
                blocks = regions_to_blocks(rows, sheet=sh.name, method="xlrd", start_table_index=table_index)
                table_index += len(blocks)
                out.blocks.extend(blocks)
                sheets_info.append({"name": sh.name, "tables": len(blocks)})
                book.unload_sheet(sheet_idx)
        finally:
            book.release_resources()
        out.page_count = len(sheets_info)
        out.metadata["sheets"] = sheets_info
        out.quality = {"cells_scanned": total_cells, "tables": table_index}
        return out


def decode_text(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace"), "latin-1"


class CsvParser(DocumentParser):
    types = frozenset({"csv"})

    def parse(self, path: Path) -> ExtractedDocument:
        settings = get_settings()
        text, encoding = decode_text(path.read_bytes())
        out = ExtractedDocument(detected_type="csv", metadata={"encoding": encoding})
        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows: list[GridRow] = []
        total = 0
        try:
            for r_idx, row in enumerate(reader, start=1):
                total += len(row)
                if total > settings.max_spreadsheet_cells:
                    raise ExtractionError("CSV exceeds the configured cell limit")
                cells = []
                for v in row:
                    v = v.strip()
                    raw: object = v
                    try:
                        if v and v.replace(",", "").replace(".", "", 1).lstrip("-").isdigit():
                            raw = float(v.replace(",", ""))
                    except ValueError:
                        pass
                    cells.append(GridCell(v, raw))
                rows.append(GridRow(number=r_idx, cells=cells))
        except csv.Error as exc:
            raise ExtractionError(f"Malformed CSV: {exc}") from exc
        formula_like = sum(1 for r in rows for c in r.cells if c.value[:1] in ("=", "+", "@") and len(c.value) > 1)
        if formula_like:
            out.warnings.append(
                f"{formula_like} cell(s) begin with formula characters; treated as plain text (CSV injection safe)."
            )
        out.blocks = regions_to_blocks(rows, sheet=None, method="csv")
        out.page_count = 1
        out.quality = {"cells_scanned": total, "tables": len(out.blocks), "delimiter": getattr(dialect, "delimiter", ",")}
        return out
