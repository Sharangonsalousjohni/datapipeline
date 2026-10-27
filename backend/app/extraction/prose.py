"""Parsers for prose formats: DOCX and TXT."""

from __future__ import annotations

import re
from pathlib import Path

from app.extraction.base import (
    HEADING,
    PARAGRAPH,
    TABLE,
    ContentBlock,
    DocumentParser,
    ExtractedDocument,
    ExtractionError,
    TableData,
)
from app.extraction.spreadsheet import decode_text
from app.extraction.tabular import render_table

_HEADING_STYLE_RE = re.compile(r"^(heading|title|subtitle)\s*(\d*)", re.I)
_TEXT_HEADING_RE = re.compile(
    r"^(#{1,6}\s+.+|(article|section|schedule|exhibit|annex|appendix|part|chapter)\s+[\w.]+.{0,80}|\d+(\.\d+)*\.?\s+[A-Z][^.!?]{0,80}|[A-Z0-9][A-Z0-9 ,&/()'\-]{3,80})$"
)


class DocxParser(DocumentParser):
    types = frozenset({"docx"})

    def parse(self, path: Path) -> ExtractedDocument:
        try:
            import docx
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError("python-docx is not installed") from exc
        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise ExtractionError(f"Could not open DOCX: {type(exc).__name__}") from exc

        out = ExtractedDocument(detected_type="docx")
        para_idx = 0
        table_idx = 0
        body = document.element.body
        for child in body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                p = Paragraph(child, document)
                text = p.text.strip()
                if not text:
                    continue
                para_idx += 1
                style = (p.style.name if p.style is not None else "") or ""
                m = _HEADING_STYLE_RE.match(style)
                level = None
                if m:
                    level = int(m.group(2)) if m.group(2) else 1
                elif len(text) < 90 and p.runs and all(r.bold for r in p.runs if r.text.strip()) and not text.endswith("."):
                    level = 3
                out.blocks.append(
                    ContentBlock(
                        block_type=HEADING if level else PARAGRAPH,
                        text=text,
                        location={"paragraph_start": para_idx, "paragraph_end": para_idx},
                        extraction_method="python-docx",
                        heading_level=level,
                    )
                )
            elif tag == "tbl":
                t = Table(child, document)
                grid: list[list[str]] = []
                for row in t.rows:
                    vals: list[str] = []
                    prev_tc = None
                    for cell in row.cells:
                        # merged cells are repeated by python-docx; keep one copy
                        if prev_tc is not None and cell._tc is prev_tc:
                            vals.append("")
                            continue
                        prev_tc = cell._tc
                        vals.append(" ".join(cell.text.split()))
                    grid.append(vals)
                grid = [r for r in grid if any(v for v in r)]
                if not grid:
                    continue
                table_idx += 1
                width = max(len(r) for r in grid)
                grid = [r + [""] * (width - len(r)) for r in grid]
                table = TableData(
                    headers=grid[0], rows=grid[1:], row_numbers=list(range(2, len(grid) + 1)), header_row_number=1
                )
                out.blocks.append(
                    ContentBlock(
                        block_type=TABLE,
                        text=render_table(table),
                        location={
                            "table_index": table_idx,
                            "paragraph_start": para_idx,
                            "paragraph_end": para_idx,
                            "row_start": 1,
                            "row_end": len(grid),
                            "col_start": 1,
                            "col_end": width,
                        },
                        extraction_method="python-docx",
                        table=table,
                    )
                )
        inline_images = len(document.inline_shapes)
        if inline_images:
            out.warnings.append(
                f"{inline_images} embedded image(s) were not OCR'd; text inside images may be missing."
            )
        out.quality = {"paragraphs": para_idx, "tables": table_idx, "embedded_images": inline_images}
        return out


class TextParser(DocumentParser):
    types = frozenset({"txt"})

    def parse(self, path: Path) -> ExtractedDocument:
        text, encoding = decode_text(path.read_bytes())
        out = ExtractedDocument(detected_type="txt", metadata={"encoding": encoding})
        para_idx = 0
        # Paragraphs are separated by blank lines; offsets refer to the decoded text.
        for m in re.finditer(r"(?:[^\n]*\S[^\n]*(?:\n|$))+", text):
            raw = m.group(0)
            stripped = raw.strip()
            if not stripped:
                continue
            start = m.start() + (len(raw) - len(raw.lstrip()))
            end = start + len(stripped)
            line_start = text.count("\n", 0, start) + 1
            line_end = line_start + stripped.count("\n")
            lines = stripped.split("\n")
            # A first line that looks like a heading becomes its own block.
            if len(lines) > 1 and _TEXT_HEADING_RE.match(lines[0].strip()) and len(lines[0]) < 90:
                para_idx += 1
                h = lines[0].strip()
                out.blocks.append(
                    ContentBlock(
                        HEADING, h.lstrip("# ").strip(),
                        {"paragraph_start": para_idx, "paragraph_end": para_idx, "char_start": start,
                         "char_end": start + len(lines[0].rstrip()), "line_start": line_start, "line_end": line_start},
                        "text", heading_level=_md_level(h),
                    )
                )
                body_start = text.index(lines[1], start + len(lines[0]))
                stripped = text[body_start:end]
                start = body_start
                line_start += 1
            para_idx += 1
            is_heading = "\n" not in stripped and len(stripped) < 90 and bool(_TEXT_HEADING_RE.match(stripped))
            out.blocks.append(
                ContentBlock(
                    HEADING if is_heading else PARAGRAPH,
                    stripped.lstrip("# ").strip() if is_heading else stripped,
                    {"paragraph_start": para_idx, "paragraph_end": para_idx, "char_start": start, "char_end": end,
                     "line_start": line_start, "line_end": line_end},
                    "text",
                    heading_level=_md_level(stripped) if is_heading else None,
                )
            )
        out.page_count = 1
        out.quality = {"paragraphs": para_idx, "chars": len(text)}
        return out


def _md_level(h: str) -> int:
    m = re.match(r"^(#{1,6})\s", h)
    if m:
        return len(m.group(1))
    m = re.match(r"^(\d+(?:\.\d+)*)", h)
    if m:
        return min(6, m.group(1).count(".") + 1)
    return 2
